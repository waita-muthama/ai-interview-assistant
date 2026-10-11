"""Parse uploaded resume into plain text with strict size/type limits."""
from io import BytesIO
from pathlib import Path
from fastapi import HTTPException
from docx import Document
from pypdf import PdfReader
from .config import MAX_DOCUMENT_CHARS

SUPPORTED = {'.pdf', '.docx', '.txt', '.md'}


def extract_document(filename: str, payload: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED:
        raise HTTPException(415, 'Only PDF, DOCX, TXT and MD resumes are accepted.')
    try:
        if extension == '.pdf':
            reader = PdfReader(BytesIO(payload))
            if reader.is_encrypted:
                raise ValueError('Password-protected PDFs are not supported.')
            text = '\n'.join(page.extract_text() or '' for page in reader.pages[:30])
        elif extension == '.docx':
            document = Document(BytesIO(payload))
            text = '\n'.join(p.text for p in document.paragraphs)
            for table in document.tables:
                text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows)
        else:
            text = payload.decode('utf-8-sig')
    except Exception as exc:
        # Intentional: provide only a safe error message, not file contents.
        if isinstance(exc, HTTPException):
            raise
        raise HTTPException(422, 'Could not read that document. Try a text-based PDF, DOCX or TXT file.') from exc
    cleaned = ' '.join(text.split())[:MAX_DOCUMENT_CHARS]
    if not cleaned:
        raise HTTPException(422, 'The document contains no extractable text. Scanned image PDFs require OCR.')
    return cleaned
